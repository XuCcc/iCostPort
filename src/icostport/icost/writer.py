"""使用 openpyxl 写出 .xlsx：首行为表头，数据自第 2 行起。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from icostport.icost.schema import ICOST_HEADERS, ICostRow


def write_icost_workbook(path: Path, rows: list[ICostRow]) -> None:
    """写入 ICost 模板兼容工作簿（单表、首行表头与 :data:`ICOST_HEADERS` 一致）。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sheet1"

    for col, title in enumerate(ICOST_HEADERS, start=1):
        ws.cell(row=1, column=col, value=title)

    for i, row in enumerate(rows, start=2):
        values = (
            row.date_display,
            row.txn_type,
            float(row.amount),
            row.primary,
            row.secondary,
            row.account1,
            row.account2,
            row.note,
            row.currency,
            row.tags,
        )
        for col, val in enumerate(values, start=1):
            ws.cell(row=i, column=col, value=val)

    wb.save(path)
