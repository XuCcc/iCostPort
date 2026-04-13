"""微信支付账单解析器。"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from icostport.core.model import Transaction
from icostport.sources.registry import register_detector, register_parser


def _normalize_cell(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _parse_datetime(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    s = _normalize_cell(value)
    if not s:
        raise ValueError("日期为空")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return datetime.fromisoformat(s.replace("Z", "+00:00"))


def _parse_amount(value: object) -> Decimal:
    if value is None:
        raise ValueError("金额为空")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    s = _normalize_cell(value)
    if not s:
        raise ValueError("金额为空")
    s = s.replace("¥", "").replace("￥", "").replace(",", "").replace(" ", "")
    try:
        return Decimal(s)
    except InvalidOperation as exc:
        raise ValueError(f"金额无效: {value!r}") from exc


def _extract_from_row(row: tuple[object, ...], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return _normalize_cell(row[index])


def parse_wechat(path: Path) -> list[Transaction]:
    """解析微信支付账单 XLSX 文件。"""
    out: list[Transaction] = []

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    header_index: int | None = None
    headers: dict[str, int] = {}
    required = {"交易时间", "收/支", "金额(元)"}

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = [_normalize_cell(cell) for cell in row]
        if not any(values):
            continue

        normalized = {cell.lower() for cell in values if cell}
        if required.issubset(normalized):
            header_index = row_idx
            for index, cell in enumerate(values):
                if cell:
                    headers[cell.strip().lower()] = index
            continue

        if header_index is None:
            continue

        if row_idx <= header_index:
            continue

        if not any(_normalize_cell(cell) for cell in row):
            continue

        raw_date = _extract_from_row(row, headers.get("交易时间"))
        raw_amount = _extract_from_row(row, headers.get("金额(元)"))
        raw_direction = _extract_from_row(row, headers.get("收/支"))

        if not raw_date or not raw_amount:
            continue

        txn_type = "支出" if "支出" in raw_direction else "收入" if "收入" in raw_direction else "中性"

        tx_type = _extract_from_row(row, headers.get("交易类型"))
        counterparty = _extract_from_row(row, headers.get("交易对方"))
        product = _extract_from_row(row, headers.get("商品"))
        remark = _extract_from_row(row, headers.get("备注"))
        payment_method = _extract_from_row(row, headers.get("支付方式"))

        note_parts = [part for part in (tx_type, counterparty, product, remark) if part and part != "/"]
        note = " ".join(note_parts).strip()

        txn = Transaction(
            occurred_at=_parse_datetime(raw_date),
            amount=_parse_amount(raw_amount),
            txn_type=txn_type,
            note=note,
            account1=payment_method,
            account2="",
            currency="CNY",
            source_path=path,
        )
        out.append(txn)

    return out


def detect_wechat_by_name(path: Path) -> bool:
    stem = path.stem.lower()
    return "微信" in stem


def detect_wechat_by_content(path: Path) -> bool:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        first_row = next(ws.iter_rows(values_only=True), ())
        return any("微信" in _normalize_cell(cell) for cell in first_row)
    except Exception:
        return False


register_detector(detect_wechat_by_name, parse_wechat, stage="filename")
register_detector(detect_wechat_by_content, parse_wechat, stage="content")
register_parser("xlsx", parse_wechat)
